"""Excel archive export with literal-text cells and bounded cell lengths."""
from __future__ import annotations

import json
from pathlib import Path
import re
import tempfile
from typing import Any

from .store import EmailStore


EXCEL_CELL_LIMIT = 32_767
_TRUNCATION_MARKER = "\n[truncated for Excel; full text remains in SQLite]"
_INVALID_XML_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

EMAIL_COLUMNS = (
    ("source_id", "Source ID"),
    ("source_name", "Source"),
    ("project_or_list", "Project / List"),
    ("thread_id", "Thread ID"),
    ("message_id", "Message ID"),
    ("in_reply_to", "In Reply To"),
    ("references", "References"),
    ("timestamp", "Timestamp"),
    ("sender", "Sender"),
    ("recipients", "Recipients"),
    ("subject", "Subject"),
    ("body_raw", "Body Raw"),
    ("current_message", "Current Authored Text"),
    ("quoted_history", "Quoted History"),
    ("attachment_names", "Attachment Filenames"),
    ("provenance", "Provenance"),
    ("license_or_terms_note", "License / Terms Note"),
    ("scope", "Scope"),
    ("labels", "Human / Source Labels"),
    ("spans", "Extraction Spans"),
    ("annotation_method", "Annotation Method"),
    ("rule_status", "Rule Baseline Status"),
    ("rule_labels", "Rule Baseline Suggestions"),
    ("rule_evidence", "Rule Baseline Evidence"),
    ("extraction_suggestions", "Rule Extraction Suggestions (Unvalidated)"),
    ("classifier_status", "Classifier Review Status"),
    ("classifier_method", "Classifier Method"),
    ("classifier_run_id", "Classifier Run ID"),
    ("classifier_labels", "Classifier Suggestions"),
    ("classifier_scores", "Classifier Scores"),
    ("body_sha256", "Body SHA-256"),
)

THREAD_COLUMNS = (
    ("thread_id", "Thread ID"),
    ("message_count", "Message Count"),
    ("first_timestamp", "First Message"),
    ("latest_timestamp", "Latest Message"),
    ("source_names", "Sources"),
    ("project_or_lists", "Projects / Lists"),
)

EXTRACTION_COLUMNS = (
    ("source_id", "Source ID"),
    ("thread_id", "Thread ID"),
    ("type", "Provisional Target"),
    ("status", "Review Status"),
    ("text", "Exact Source Text"),
    ("field", "Source Field"),
    ("start", "Start Offset"),
    ("end", "End Offset"),
    ("rule_id", "Rule ID"),
    ("human_gold", "Human Gold"),
)


def _display_value(value: Any) -> str | int | float | bool | None:
    if value is None:
        return None
    if isinstance(value, (str, int, float, bool)):
        return _INVALID_XML_CONTROL.sub("\ufffd", value) if isinstance(value, str) else value
    return _INVALID_XML_CONTROL.sub(
        "\ufffd", json.dumps(value, ensure_ascii=False, sort_keys=True)
    )


def _truncate_text(value: str) -> tuple[str, bool]:
    if len(value) <= EXCEL_CELL_LIMIT:
        return value, False
    keep = EXCEL_CELL_LIMIT - len(_TRUNCATION_MARKER)
    return value[:keep] + _TRUNCATION_MARKER, True


def _append_literal_row(sheet: Any, values: list[Any]) -> None:
    from openpyxl.cell.cell import WriteOnlyCell

    cells = []
    for value in values:
        cell = WriteOnlyCell(sheet, value=None)
        if isinstance(value, str):
            text, _ = _truncate_text(value)
            cell.value = text
            # Keep formula-looking mail text as an inline string cell.
            cell.data_type = "s"
        else:
            cell.value = value
        cells.append(cell)
    sheet.append(cells)


def export_excel(store: EmailStore, path: str | Path) -> Path:
    """Export messages, threads, and unvalidated extraction suggestions atomically."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("Excel export needs openpyxl; install desktop/requirements.txt.") from exc

    target = Path(path).expanduser().resolve()
    if target.suffix.lower() != ".xlsx":
        raise ValueError("Excel export path must end in .xlsx")
    target.parent.mkdir(parents=True, exist_ok=True)
    emails = store.list_emails()
    threads = store.list_threads()
    workbook = Workbook(write_only=True)
    email_sheet = workbook.create_sheet("Emails")
    thread_sheet = workbook.create_sheet("Threads")
    extraction_sheet = workbook.create_sheet("Extraction Suggestions")
    header_fill = PatternFill("solid", fgColor="174A5B")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(vertical="center", wrap_text=True)

    for sheet, columns in (
        (email_sheet, EMAIL_COLUMNS),
        (thread_sheet, THREAD_COLUMNS),
        (extraction_sheet, EXTRACTION_COLUMNS),
    ):
        sheet.freeze_panes = "A2"
        header_cells = []
        for _key, label in columns:
            from openpyxl.cell.cell import WriteOnlyCell

            cell = WriteOnlyCell(sheet, value=label)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            header_cells.append(cell)
        sheet.append(header_cells)
        count = (
            len(emails)
            if sheet is email_sheet
            else len(threads)
            if sheet is thread_sheet
            else max(1, sum(len(row.get("extraction_suggestions", {}).get("spans", [])) for row in emails))
        )
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, count) + 1}"

    for row in emails:
        values = [_display_value(row.get(key)) for key, _label in EMAIL_COLUMNS]
        _append_literal_row(email_sheet, values)
    for row in threads:
        values = [_display_value(row.get(key)) for key, _label in THREAD_COLUMNS]
        _append_literal_row(thread_sheet, values)
    for row in emails:
        extraction = row.get("extraction_suggestions", {})
        for span in extraction.get("spans", []):
            values = {
                **span,
                "source_id": row["source_id"],
                "thread_id": row["thread_id"],
                "human_gold": False,
            }
            _append_literal_row(
                extraction_sheet,
                [_display_value(values.get(key)) for key, _label in EXTRACTION_COLUMNS],
            )

    handle = tempfile.NamedTemporaryFile(
        prefix=f".{target.stem}.", suffix=".tmp.xlsx", dir=target.parent, delete=False
    )
    temp_path = Path(handle.name)
    handle.close()
    try:
        workbook.save(temp_path)
        temp_path.replace(target)
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise
    return target

"""Reviewed project facts and immutable history exported as literal Excel cells."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import tempfile

from .excel_export import _append_literal_row, _display_value
from .item_contracts import FIELDS, SCHEMA


def project_workbook_bytes(ledger) -> bytes:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    snapshot = ledger.export_snapshot()
    workbook = Workbook(write_only=True)
    items = workbook.create_sheet("Items")
    history = workbook.create_sheet("History")
    metadata = workbook.create_sheet("Metadata")
    item_columns = ("item_id", "version", *FIELDS, "source_id", "input_fingerprint", "source_revision", "source_conflict")
    history_columns = ("event_id", "item_id", "action", "before", "after", "note", "source_id", "input_fingerprint", "source_revision", "candidate_id", "timestamp", "evidence")
    for sheet, columns in ((items, item_columns), (history, history_columns)):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
        _append_literal_row(sheet, list(columns))
    for item in snapshot:
        values = {**item, **item["fields"]}
        _append_literal_row(items, [_display_value(values.get(k)) for k in item_columns])
        for event in item.get("history", []):
            _append_literal_row(history, [_display_value(event.get(k)) for k in history_columns])
    _append_literal_row(metadata, ["product_contract", SCHEMA])
    _append_literal_row(metadata, ["meaning", "Approved facts and audited changes; deterministic suggestions are not human gold. Full source/history remains in SQLite."])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def export_project_excel(ledger, path: str | Path) -> Path:
    target = Path(path).expanduser().resolve()
    if target.suffix.lower() != ".xlsx":
        raise ValueError("Export path must end in .xlsx")
    data = project_workbook_bytes(ledger)
    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=target.parent, suffix=".tmp.xlsx", delete=False) as handle:
        temporary = Path(handle.name)
        handle.write(data)
    try:
        temporary.replace(target)
    finally:
        temporary.unlink(missing_ok=True)
    return target

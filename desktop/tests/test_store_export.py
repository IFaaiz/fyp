from __future__ import annotations

from contextlib import closing
from pathlib import Path
import sqlite3
import unittest

from desktop.local_product.adapters import EmlAdapter
from desktop.local_product.excel_export import export_excel
from desktop.local_product.models import EmailRecord
from desktop.local_product.store import AnnotationConflictError, EmailStore
from desktop.tests.support import temporary_directory


def _record(
    source_id: str,
    *,
    body: str = "The report is attached.",
    thread_id: str = "thread-1",
    subject: str = "Project update",
    message_id: str | None = None,
    in_reply_to: str | None = None,
    references: tuple[str, ...] = (),
    timestamp: str = "2026-10-09T09:30:00+00:00",
    annotation_method: str = "UNANNOTATED",
    labels: tuple[str, ...] = (),
) -> EmailRecord:
    return EmailRecord(
        source_id=source_id,
        source_name="fixture",
        project_or_list="project-list",
        message_id=message_id,
        in_reply_to=in_reply_to,
        references=references,
        thread_id=thread_id,
        subject=subject,
        sender="sender@example.test",
        recipients=("team@example.test",),
        timestamp=timestamp,
        body_raw=body,
        current_message=body,
        provenance={"source": "synthetic"},
        labels=labels,
        annotation_method=annotation_method,
    )


class StoreTests(unittest.TestCase):
    def test_schema_v1_database_migrates_to_extraction_suggestions_column(self):
        with temporary_directory() as temp:
            path = Path(temp) / "archive.sqlite3"
            store = EmailStore(path)
            with closing(sqlite3.connect(path)) as connection, connection:
                connection.execute("ALTER TABLE emails DROP COLUMN extraction_suggestions_json")
                connection.execute("PRAGMA user_version = 1")
            migrated = EmailStore(path)
            with closing(sqlite3.connect(path)) as connection:
                columns = {row[1] for row in connection.execute("PRAGMA table_info(emails)")}
                version = connection.execute("PRAGMA user_version").fetchone()[0]
            self.assertIn("extraction_suggestions_json", columns)
            self.assertEqual(version, 2)

    def test_duplicate_import_is_idempotent_and_thread_relationships_are_preserved(self):
        with temporary_directory() as temp:
            store = EmailStore(Path(temp) / "archive.sqlite3")
            rows = [
                _record("root", message_id="root@example.test"),
                _record(
                    "reply",
                    message_id="reply@example.test",
                    in_reply_to="root@example.test",
                    references=("root@example.test",),
                    timestamp="2026-10-10T09:30:00+00:00",
                ),
            ]
            first = store.import_records(rows)
            second = store.import_records(rows)
            self.assertEqual(first, {"inserted": 2, "updated": 0, "total": 2})
            self.assertEqual(second, {"inserted": 0, "updated": 2, "total": 2})
            self.assertEqual(store.count(), 2)
            thread = store.list_threads()[0]
            self.assertEqual(thread["thread_id"], "thread-1")
            self.assertEqual(thread["message_count"], 2)
            reply = store.get_email("reply")
            self.assertEqual(reply["in_reply_to"], "root@example.test")
            self.assertEqual(reply["references"], ["root@example.test"])
            self.assertEqual(reply["rule_status"], "ABSTAIN")
            self.assertFalse(reply["labels"])
            self.assertEqual(reply["annotation_method"], "UNANNOTATED")
            self.assertEqual(reply["extraction_suggestions"]["human_gold"], False)
            self.assertEqual(reply["extraction_suggestions"]["status"], "ABSTAIN")

    def test_message_ids_reconstruct_parent_reply_thread_without_references(self):
        with temporary_directory() as temp:
            store = EmailStore(Path(temp) / "archive.sqlite3")
            parent = _record("parent", thread_id="", message_id="root@example.test", references=())
            reply = _record(
                "reply",
                thread_id="",
                message_id="reply@example.test",
                in_reply_to="root@example.test",
                references=(),
            )
            store.import_records([reply, parent])
            saved_parent = store.get_email("parent")
            saved_reply = store.get_email("reply")
            self.assertEqual(saved_parent["thread_id"], "message-thread:root@example.test")
            self.assertEqual(saved_reply["thread_id"], saved_parent["thread_id"])
            self.assertEqual(store.list_threads()[0]["message_count"], 2)

    def test_changed_authored_text_reimport_fails_closed_and_preserves_annotations(self):
        with temporary_directory() as temp:
            store = EmailStore(Path(temp) / "archive.sqlite3")
            record = _record("owned", body="First source body.")
            store.import_records([record])
            with closing(sqlite3.connect(store.path)) as connection, connection:
                connection.execute(
                    "UPDATE emails SET scope=?, labels_json=?, annotation_method=? WHERE source_id=?",
                    ("PROJECT", '["APPROVAL"]', "HUMAN_REVIEWED", "owned"),
                )
            for refreshed in (
                _record("owned", body="Corrected source body."),
                _record("owned", body="First source body.", subject="Changed subject"),
            ):
                with self.assertRaises(AnnotationConflictError):
                    store.import_records([refreshed])
            saved = store.get_email("owned")
            self.assertEqual(saved["body_raw"], "First source body.")
            self.assertEqual(saved["current_message"], "First source body.")
            self.assertEqual(saved["scope"], "PROJECT")
            self.assertEqual(saved["labels"], ["APPROVAL"])
            self.assertEqual(saved["annotation_method"], "HUMAN_REVIEWED")

    def test_import_batch_rolls_back_when_record_is_not_normalized(self):
        with temporary_directory() as temp:
            store = EmailStore(Path(temp) / "archive.sqlite3")

            def records():
                yield _record("would-have-been-inserted")
                yield object()

            with self.assertRaises(TypeError):
                store.import_records(records())
            self.assertEqual(store.count(), 0)
            self.assertEqual(store.list_threads(), [])

    def test_rule_baseline_and_existing_classifier_fields_are_separate_from_labels(self):
        with temporary_directory() as temp:
            store = EmailStore(Path(temp) / "archive.sqlite3")
            record = _record("suggest", body="Please send the revised report by Friday.")
            store.import_records([record])
            saved = store.get_email("suggest")
            self.assertEqual(set(saved["rule_labels"]), {"REPORT_REQUEST", "DEADLINE"})
            self.assertEqual(saved["rule_status"], "REVIEW")
            self.assertEqual(saved["labels"], [])
            self.assertEqual(saved["annotation_method"], "UNANNOTATED")

    def test_import_stores_rule_extraction_separately_without_promoting_spans(self):
        with temporary_directory() as temp:
            store = EmailStore(Path(temp) / "archive.sqlite3")
            store.import_records([_record("extract", body="Please complete the report by Friday.")])
            saved = store.get_email("extract")
            suggestion = saved["extraction_suggestions"]
            self.assertEqual(suggestion["status"], "REVIEW")
            self.assertFalse(suggestion["human_gold"])
            self.assertTrue(suggestion["spans"])
            self.assertEqual(saved["labels"], [])
            self.assertEqual(saved["spans"], [])
            self.assertEqual(saved["annotation_method"], "UNANNOTATED")
            self.assertIsNone(saved["classifier_status"])


class ExcelExportTests(unittest.TestCase):
    def test_formula_looking_email_text_exports_as_literal_and_threads_are_included(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")
        with temporary_directory() as temp:
            root = Path(temp)
            store = EmailStore(root / "archive.sqlite3")
            formula_subject = '=HYPERLINK("https://example.test","click")'
            formula_body = "+SUM(1,1)\n@mention"
            record = EmailRecord(
                source_id="formula-source",
                source_name="fixture",
                thread_id="formula-thread",
                subject=formula_subject,
                sender="sender@example.test",
                timestamp="2026-10-09T09:30:00Z",
                body_raw=formula_body,
                current_message=formula_body,
                annotation_method="UNANNOTATED",
            )
            store.import_records([record])
            destination = export_excel(store, root / "archive.xlsx")
            workbook = load_workbook(destination, data_only=False, read_only=False)
            self.assertEqual(workbook.sheetnames, ["Emails", "Threads", "Extraction Suggestions"])
            sheet = workbook["Emails"]
            headers = {cell.value: cell.column for cell in sheet[1]}
            subject_cell = sheet.cell(row=2, column=headers["Subject"])
            body_cell = sheet.cell(row=2, column=headers["Body Raw"])
            self.assertEqual(subject_cell.value, formula_subject)
            self.assertEqual(subject_cell.data_type, "s")
            self.assertEqual(body_cell.value, formula_body)
            self.assertEqual(body_cell.data_type, "s")
            self.assertEqual(workbook["Threads"].cell(row=2, column=2).value, 1)
            workbook.close()

    def test_extraction_sheet_and_json_column_are_unvalidated_review_only(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")
        with temporary_directory() as temp:
            root = Path(temp)
            store = EmailStore(root / "archive.sqlite3")
            store.import_records([_record("extract", body="Please complete the report by Friday.")])
            path = export_excel(store, root / "archive.xlsx")
            workbook = load_workbook(path, data_only=True, read_only=True)
            emails = workbook["Emails"]
            headers = {cell.value: cell.column for cell in next(emails.iter_rows(min_row=1, max_row=1))}
            row = next(emails.iter_rows(min_row=2, values_only=True))
            value = row[headers["Rule Extraction Suggestions (Unvalidated)"] - 1]
            try:
                self.assertIn('"human_gold": false', value)
                suggestions = workbook["Extraction Suggestions"]
                suggestion_headers = [cell.value for cell in next(suggestions.iter_rows(min_row=1, max_row=1))]
                self.assertIn("Provisional Target", suggestion_headers)
                suggestion_row = next(suggestions.iter_rows(min_row=2, values_only=True))
                self.assertIn(suggestion_row[suggestion_headers.index("Review Status")], {"REVIEW", "ABSTAIN"})
                self.assertFalse(suggestion_row[suggestion_headers.index("Human Gold")])
            finally:
                workbook.close()

    def test_oversize_excel_cell_is_marked_as_truncated(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl is not installed")
        with temporary_directory() as temp:
            root = Path(temp)
            store = EmailStore(root / "archive.sqlite3")
            body = "x" * 33_000
            store.import_records([_record("long", body=body)])
            path = export_excel(store, root / "archive.xlsx")
            workbook = load_workbook(path, read_only=True)
            sheet = workbook["Emails"]
            headers = {cell.value: cell.column for cell in next(sheet.iter_rows(min_row=1, max_row=1))}
            exported = next(sheet.iter_rows(min_row=2, values_only=True))
            value = exported[headers["Body Raw"] - 1]
            self.assertLessEqual(len(value), 32_767)
            self.assertIn("full text remains in SQLite", value)
            workbook.close()


if __name__ == "__main__":
    unittest.main()

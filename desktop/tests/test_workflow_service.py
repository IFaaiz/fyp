"""HTTP-only synthetic coverage for the reviewed project-item workflow."""
from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from io import BytesIO
import json
from http.client import HTTPConnection
from pathlib import Path
from threading import Barrier
import unittest
from urllib.error import HTTPError
from urllib.request import Request, urlopen
from urllib.parse import urlsplit

from desktop.local_product.backup import backup_database
from desktop.local_product.models import EmailRecord
from desktop.local_product.service import MAX_REQUEST_BYTES, start_local_service
from desktop.local_product.store import EmailStore
from desktop.tests.support import temporary_directory


TOKEN = "synthetic-owner-token-6f4a"


def _record(
    source_id: str,
    body: str,
    *,
    message_id: str,
    in_reply_to: str | None = None,
) -> EmailRecord:
    return EmailRecord(
        source_id=source_id,
        source_name="synthetic",
        message_id=message_id,
        in_reply_to=in_reply_to,
        references=("root@example.test",) if in_reply_to else (),
        thread_id="thread-1",
        subject="Synthetic project update",
        sender="sender@example.test",
        recipients=("reviewers@example.test",),
        timestamp="2026-10-11T06:00:00Z",
        body_raw=body,
        current_message=body,
        provenance={"adapter": "synthetic", "source_access_scope": "fixture-only"},
    )


def _request(
    base_url: str,
    method: str,
    path: str,
    *,
    data=None,
    token: str | None = TOKEN,
    headers: dict[str, str] | None = None,
    raw_body: bytes | None = None,
) -> tuple[int, dict[str, str], bytes]:
    request_headers = dict(headers or {})
    if token is not None:
        request_headers["Authorization"] = f"Bearer {token}"
    payload = raw_body
    if raw_body is None and data is not None:
        payload = json.dumps(data, ensure_ascii=False).encode("utf-8")
    if payload is not None:
        request_headers.setdefault("Content-Type", "application/json")
    request = Request(base_url + path, data=payload, method=method, headers=request_headers)
    try:
        with urlopen(request, timeout=10) as response:
            return response.status, dict(response.headers.items()), response.read()
    except HTTPError as error:
        return error.code, dict(error.headers.items()), error.read()


def _json_request(base_url: str, method: str, path: str, **kwargs):
    status, headers, body = _request(base_url, method, path, **kwargs)
    return status, headers, json.loads(body.decode("utf-8"))


def _oversized_length_request(base_url: str) -> tuple[int, dict[str, str], bytes]:
    parsed = urlsplit(base_url)
    host = parsed.hostname or "127.0.0.1"
    port = parsed.port
    if port is None:
        raise AssertionError("Local service URL is missing a port")
    connection = HTTPConnection(host, port, timeout=10)
    try:
        connection.putrequest("POST", "/api/suggest", skip_host=True)
        connection.putheader("Host", f"{host}:{port}")
        connection.putheader("Authorization", f"Bearer {TOKEN}")
        connection.putheader("Content-Type", "application/json")
        connection.putheader("Content-Length", str(MAX_REQUEST_BYTES + 1))
        connection.putheader("Connection", "close")
        connection.endheaders()
        response = connection.getresponse()
        return response.status, dict(response.headers.items()), response.read()
    finally:
        connection.close()


class WorkflowServiceTests(unittest.TestCase):
    def setUp(self):
        self._scratch_context = temporary_directory()
        self.temp_path = Path(self._scratch_context.__enter__())
        self.addCleanup(self._scratch_context.__exit__, None, None, None)
        self.root_text = "Please review the weekly report by Friday. Please send the revised agenda by Monday."
        self.reply_text = (
            "Please schedule a meeting with the department on Tuesday at 3 pm. "
            "Please review and approve the minutes."
        )
        self.store = EmailStore(self.temp_path / "archive.sqlite3")
        self.store.import_records(
            [
                _record("root", self.root_text, message_id="root@example.test"),
                _record(
                    "reply", self.reply_text, message_id="reply@example.test",
                    in_reply_to="root@example.test",
                ),
            ]
        )
        self.service = start_local_service(self.store, TOKEN)
        self.addCleanup(self.service.close)

    def test_authorized_http_workflow_export_and_backup_restore(self):
        status, _headers, emails = _json_request(self.service.base_url, "GET", "/api/emails")
        self.assertEqual(status, 200)
        self.assertEqual(len(emails["emails"]), 2)
        self.assertEqual(len({row["thread_id"] for row in emails["emails"]}), 1)
        self.assertEqual(_json_request(self.service.base_url, "GET", "/api/candidates")[2]["candidates"], [])
        self.assertEqual(_json_request(self.service.base_url, "GET", "/api/items")[2]["items"], [])

        root_suggestions = _json_request(
            self.service.base_url, "POST", "/api/suggest", data={"source_id": "root"}
        )[2]["candidates"]
        reply_suggestions = _json_request(
            self.service.base_url, "POST", "/api/suggest", data={"source_id": "reply"}
        )[2]["candidates"]
        report = next(item for item in root_suggestions if "weekly report" in item["proposed_fields"]["title"])
        agenda = next(item for item in root_suggestions if "revised agenda" in item["proposed_fields"]["title"])
        meeting = next(item for item in reply_suggestions if item["proposed_fields"]["kind"] == "MEETING")
        approval = next(item for item in reply_suggestions if item["proposed_fields"]["kind"] == "APPROVAL")

        report_decision = {
            "candidate_id": report["candidate_id"],
            "decision": "APPROVE",
            "expected_version": report["version"],
            "note": "Verified against the source sentence.",
        }
        status, _headers, report_item = _json_request(
            self.service.base_url, "POST", "/api/decide", data=report_decision
        )
        self.assertEqual(status, 200)
        self.assertEqual(report_item["fields"]["kind"], "DOCUMENT_REQUEST")
        self.assertEqual(report_item["fields"]["status"], "OPEN")
        report_id = report_item["item_id"]

        # An identical replay is idempotent; a conflicting replay is rejected.
        replay_status, _headers, replay = _json_request(
            self.service.base_url, "POST", "/api/decide", data=report_decision
        )
        self.assertEqual(replay_status, 200)
        self.assertEqual(replay["item_id"], report_id)
        conflict_status, _headers, conflict = _json_request(
            self.service.base_url,
            "POST",
            "/api/decide",
            data={**report_decision, "decision": "REJECT"},
        )
        self.assertEqual(conflict_status, 409)
        self.assertEqual(set(conflict), {"error"})

        rejected_status, _headers, rejected = _json_request(
            self.service.base_url,
            "POST",
            "/api/decide",
            data={
                "candidate_id": agenda["candidate_id"],
                "decision": "REJECT",
                "expected_version": agenda["version"],
                "note": "Duplicate request; retain the approved report item only.",
            },
        )
        self.assertEqual(rejected_status, 200, rejected)
        self.assertEqual(rejected["status"], "REJECTED")

        corrected_status, _headers, corrected = _json_request(
            self.service.base_url,
            "POST",
            "/api/decide",
            data={
                "candidate_id": meeting["candidate_id"],
                "decision": "CORRECT",
                "expected_version": meeting["version"],
                "fields": {"owner": "Operations", "due": "2026-10-13", "due_text": "Tuesday"},
                "note": "Set the accountable team and confirmed meeting date.",
            },
        )
        self.assertEqual(corrected_status, 200)
        self.assertEqual(corrected["fields"]["owner"], "Operations")
        self.assertEqual(corrected["fields"]["due"], "2026-10-13")

        update_fields = {"owner": "Analyst", "due": "2026-10-16", "status": "DELIVERED"}
        update_candidate = _json_request(
            self.service.base_url,
            "POST",
            "/api/propose",
            data={
                "source_id": "root",
                "target_item_id": report_id,
                "fields": update_fields,
                "evidence": report["evidence"],
            },
        )[2]
        self.assertEqual(update_candidate["target_version"], 1)
        updated_status, _headers, updated = _json_request(
            self.service.base_url,
            "POST",
            "/api/decide",
            data={
                "candidate_id": update_candidate["candidate_id"],
                "decision": "APPROVE",
                "expected_version": update_candidate["version"],
            },
        )
        self.assertEqual(updated_status, 200)
        self.assertEqual(updated["version"], 2)
        self.assertEqual(updated["fields"]["owner"], "Analyst")
        self.assertEqual(updated["fields"]["due"], "2026-10-16")
        self.assertEqual(updated["fields"]["status"], "DELIVERED")

        competing = []
        for owner, due in (("Manager", "2026-10-17"), ("Director", "2026-10-18")):
            competing.append(
                _json_request(
                    self.service.base_url,
                    "POST",
                    "/api/propose",
                    data={
                        "source_id": "root",
                        "target_item_id": report_id,
                        "fields": {"owner": owner, "due": due},
                        "evidence": report["evidence"],
                    },
                )[2]
            )
        first_update_status, _headers, after_competitor = _json_request(
            self.service.base_url,
            "POST",
            "/api/decide",
            data={
                "candidate_id": competing[0]["candidate_id"],
                "decision": "APPROVE",
                "expected_version": competing[0]["version"],
            },
        )
        self.assertEqual(first_update_status, 200)
        second_update_status, _headers, stale_update = _json_request(
            self.service.base_url,
            "POST",
            "/api/decide",
            data={
                "candidate_id": competing[1]["candidate_id"],
                "decision": "APPROVE",
                "expected_version": competing[1]["version"],
            },
        )
        self.assertEqual(second_update_status, 409)
        self.assertEqual(set(stale_update), {"error"})

        undo_status, _headers, undone = _json_request(
            self.service.base_url,
            "POST",
            "/api/undo",
            data={"item_id": report_id, "expected_version": after_competitor["version"], "note": "Undo competing assignment."},
        )
        self.assertEqual(undo_status, 200)
        self.assertEqual(undone["version"], after_competitor["version"] + 1)
        self.assertEqual(undone["fields"]["owner"], "Analyst")
        self.assertEqual(undone["fields"]["status"], "DELIVERED")

        # Two independent candidates based on the same target version race over
        # HTTP. Only one may commit, and the losing decision must leave no event.
        concurrent_candidates = []
        for owner, due in (("Planner A", "2026-10-19"), ("Planner B", "2026-10-20")):
            concurrent_candidates.append(
                _json_request(
                    self.service.base_url,
                    "POST",
                    "/api/propose",
                    data={
                        "source_id": "root",
                        "target_item_id": report_id,
                        "fields": {"owner": owner, "due": due},
                        "evidence": report["evidence"],
                    },
                )[2]
            )
        start_barrier = Barrier(2)

        def approve_after_barrier(candidate):
            start_barrier.wait(timeout=5)
            return _json_request(
                self.service.base_url,
                "POST",
                "/api/decide",
                data={
                    "candidate_id": candidate["candidate_id"],
                    "decision": "APPROVE",
                    "expected_version": candidate["version"],
                },
            )

        with ThreadPoolExecutor(max_workers=2) as pool:
            concurrent_results = list(pool.map(approve_after_barrier, concurrent_candidates))
        self.assertEqual(sorted(result[0] for result in concurrent_results), [200, 409])
        concurrent_item_status, _headers, concurrent_item = _json_request(
            self.service.base_url, "GET", f"/api/items/{report_id}"
        )
        self.assertEqual(concurrent_item_status, 200)
        self.assertEqual(concurrent_item["version"], undone["version"] + 1)
        self.assertEqual(len(concurrent_item["history"]), 5)

        formula_title = '=HYPERLINK("https://example.test","Open")'
        formula_note = "=1+1"
        formula_candidate = _json_request(
            self.service.base_url,
            "POST",
            "/api/propose",
            data={
                "source_id": "root",
                "fields": {"title": formula_title, "kind": "TASK"},
                "evidence": report["evidence"],
            },
        )[2]
        formula_item = _json_request(
            self.service.base_url,
            "POST",
            "/api/decide",
            data={
                "candidate_id": formula_candidate["candidate_id"],
                "decision": "APPROVE",
                "expected_version": formula_candidate["version"],
                "note": formula_note,
            },
        )[2]

        stale_candidate = _json_request(
            self.service.base_url,
            "POST",
            "/api/propose",
            data={
                "source_id": "reply",
                "fields": {"title": "Prepare the dashboard", "kind": "TASK"},
                "evidence": approval["evidence"],
            },
        )[2]
        revised_reply = replace(
            _record(
                "reply", self.reply_text + " Please prepare the dashboard.",
                message_id="reply@example.test", in_reply_to="root@example.test",
            ),
            current_message=self.reply_text + " Please prepare the dashboard.",
        )
        self.store.import_records([revised_reply])
        stale_status, _headers, stale_response = _json_request(
            self.service.base_url,
            "POST",
            "/api/decide",
            data={
                "candidate_id": stale_candidate["candidate_id"],
                "decision": "APPROVE",
                "expected_version": stale_candidate["version"],
            },
        )
        self.assertEqual(stale_status, 409)
        self.assertEqual(set(stale_response), {"error"})
        stale_replay_status, _headers, stale_replay = _json_request(
            self.service.base_url,
            "POST",
            "/api/decide",
            data={
                "candidate_id": stale_candidate["candidate_id"],
                "decision": "APPROVE",
                "expected_version": stale_candidate["version"],
            },
        )
        self.assertEqual(stale_replay_status, 409)
        self.assertEqual(set(stale_replay), {"error"})
        candidates = _json_request(self.service.base_url, "GET", "/api/candidates")[2]["candidates"]
        saved_stale = next(c for c in candidates if c["candidate_id"] == stale_candidate["candidate_id"])
        self.assertEqual(saved_stale["status"], "STALE")

        export_status, export_headers, workbook_bytes = _request(
            self.service.base_url, "POST", "/api/export", data={}
        )
        self.assertEqual(export_status, 200)
        self.assertIn("spreadsheetml.sheet", export_headers.get("Content-Type", ""))
        self.assertTrue(workbook_bytes.startswith(b"PK"))
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl is required for workbook-content assertions")
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=False, read_only=True)
        self.assertEqual(set(workbook.sheetnames), {"Items", "History", "Metadata"})
        item_sheet = workbook["Items"]
        headers = [cell.value for cell in next(item_sheet.iter_rows(min_row=1, max_row=1))]
        title_column = headers.index("title")
        formula_row = next(
            row for row in item_sheet.iter_rows(min_row=2)
            if row[headers.index("item_id")].value == formula_item["item_id"]
        )
        self.assertEqual(formula_row[title_column].value, formula_title)
        self.assertEqual(formula_row[title_column].data_type, "s")
        history_sheet = workbook["History"]
        history_headers = [cell.value for cell in next(history_sheet.iter_rows(min_row=1, max_row=1))]
        note_column = history_headers.index("note")
        formula_history = next(
            row for row in history_sheet.iter_rows(min_row=2)
            if row[history_headers.index("item_id")].value == formula_item["item_id"]
        )
        self.assertEqual(formula_history[note_column].value, formula_note)
        self.assertEqual(formula_history[note_column].data_type, "s")
        workbook.close()

        backup_path = self.temp_path / "archive-backup.sqlite3"
        backup_database(self.store.path, backup_path)
        restored_store = EmailStore(backup_path)
        with start_local_service(restored_store, TOKEN) as restored_service:
            restored_items = _json_request(restored_service.base_url, "GET", "/api/items")[2]["items"]
            restored_report_status, _headers, restored_report = _json_request(
                restored_service.base_url, "GET", f"/api/items/{report_id}"
            )
            restored_formula_status, _headers, restored_formula = _json_request(
                restored_service.base_url, "GET", f"/api/items/{formula_item['item_id']}"
            )
        self.assertGreaterEqual(len(restored_items), 3)
        self.assertEqual(restored_report_status, 200)
        self.assertEqual(restored_report["version"], concurrent_item["version"])
        self.assertEqual(len(restored_report["history"]), 5)
        self.assertEqual(restored_formula_status, 200)
        self.assertEqual(restored_formula["fields"]["title"], formula_title)

    def test_http_auth_origin_host_size_and_safe_error_contracts(self):
        protected = [
            ("GET", "/api/emails", None),
            ("GET", "/api/candidates", None),
            ("GET", "/api/items", None),
            ("GET", "/api/items/missing", None),
            ("POST", "/api/suggest", {"source_id": "root"}),
            ("POST", "/api/propose", {}),
            ("POST", "/api/decide", {}),
            ("POST", "/api/undo", {}),
            ("POST", "/api/export", {}),
        ]
        for method, path, data in protected:
            with self.subTest(method=method, path=path):
                status, headers, body = _request(
                    self.service.base_url, method, path, data=data, token=None
                )
                self.assertEqual(status, 401)
                self.assertIn("application/json", headers.get("Content-Type", ""))
                self.assertEqual(set(json.loads(body)), {"error"})

        origin_status, _headers, origin_error = _json_request(
            self.service.base_url,
            "GET",
            "/api/emails",
            headers={"Origin": "http://untrusted.example"},
        )
        self.assertEqual(origin_status, 403)
        self.assertEqual(set(origin_error), {"error"})
        host_status, _headers, host_error = _json_request(
            self.service.base_url,
            "GET",
            "/api/emails",
            headers={"Host": "untrusted.example"},
        )
        self.assertEqual(host_status, 403)
        self.assertEqual(set(host_error), {"error"})

        missing_status, _headers, missing = _json_request(
            self.service.base_url, "GET", "/api/items/not-an-item"
        )
        self.assertEqual(missing_status, 404)
        self.assertEqual(missing, {"error": "Item not found"})
        invalid_status, _headers, invalid = _json_request(
            self.service.base_url,
            "POST",
            "/api/propose",
            data={
                "source_id": "root",
                "fields": {"title": "PRIVATE_SENTINEL", "due": "next Friday"},
                "evidence": {
                    "field": "current_message",
                    "start": 0,
                    "end": len(self.root_text.split(".")[0]) + 1,
                    "text": self.root_text.split(".")[0] + ".",
                },
            },
        )
        self.assertEqual(invalid_status, 400)
        self.assertNotIn("PRIVATE_SENTINEL", json.dumps(invalid))
        self.assertEqual(set(invalid), {"error"})

        unknown_field_status, _headers, unknown_field = _json_request(
            self.service.base_url,
            "POST",
            "/api/suggest",
            data={"source_id": "root", "unknown_field": "PRIVATE_SENTINEL"},
        )
        self.assertEqual(unknown_field_status, 400)
        self.assertNotIn("PRIVATE_SENTINEL", json.dumps(unknown_field))
        self.assertEqual(set(unknown_field), {"error"})

        duplicate_json_status, _headers, duplicate_json = _json_request(
            self.service.base_url,
            "POST",
            "/api/suggest",
            raw_body=b'{"source_id":"root","source_id":"PRIVATE_SENTINEL"}',
        )
        self.assertEqual(duplicate_json_status, 400)
        self.assertNotIn("PRIVATE_SENTINEL", json.dumps(duplicate_json))

        large_status, _headers, large_body = _oversized_length_request(self.service.base_url)
        large_error = json.loads(large_body.decode("utf-8"))
        self.assertEqual(large_status, 400)
        self.assertEqual(set(large_error), {"error"})

        wrong_content_type_status, _headers, wrong_content_type = _json_request(
            self.service.base_url,
            "POST",
            "/api/suggest",
            data={"source_id": "root"},
            headers={"Content-Type": "text/plain"},
        )
        self.assertEqual(wrong_content_type_status, 400)
        self.assertEqual(set(wrong_content_type), {"error"})


if __name__ == "__main__":
    unittest.main()
